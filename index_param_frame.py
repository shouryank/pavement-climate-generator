"""
Index Parameter Calculator - Tkinter Frontend Frame
===================================================
Provides the UI for selecting a target location, scenarios, models, year ranges,
and index parameters. Runs calculations in a background thread with
a progress indicator, then displays summary popup and saves CSVs.
"""

import csv
import math
import shutil
import threading
from pathlib import Path
from threading import Thread

import tkinter as tk
from tkinter import filedialog, messagebox, ttk

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, PatternFill, Font, Border, Side

from generator_frontend import list_subdirs
from index_parameter_calculator import (
    PARAMETER_METADATA_BY_KEY,
    PARAMETER_REGISTRY,
    PG_RELIABILITY_MAP,
    PG_PARAMETER_KEYS,
    run_index_calculation,
)
from main_v3 import nearest_grid_points_from_file, resolve_paths_for_year
from multiselect_dropdown import MultiSelectDropdown
from operation_control import OperationCancelledError


PRECIP_PARAM_KEYS = {
    key for key, meta in PARAMETER_METADATA_BY_KEY.items()
    if meta.get("parameter_type") == "precip"
}

BASELINE_YEAR_MIN = 1950
BASELINE_YEAR_MAX = 2005
FUTURE_YEAR_MIN = 2006
FUTURE_YEAR_MAX = 2099
HISTORICAL_SCENARIO = "historical"


class IndexParameterFrame(ttk.Frame):
    """Index Parameter Calculator UI."""

    def __init__(self, parent, controller):
        super().__init__(parent, padding=16)
        self.c = controller

        top = ttk.Frame(self)
        top.grid(row=0, column=0, columnspan=10, sticky="ew", pady=(0, 8))
        self.back_btn = ttk.Button(top, text="Back", command=lambda: self.c.show_frame("MainMenu"))
        self.back_btn.pack(side="left")
        ttk.Label(top, text="Index Parameter Calculator", font=("Segoe UI", 14, "bold")).pack(side="left", padx=12)

        ttk.Label(self, text="Latitude:").grid(row=1, column=0, sticky="w")
        self.latitude_var = tk.StringVar()
        self.latitude_entry = ttk.Entry(self, textvariable=self.latitude_var, width=18)
        self.latitude_entry.grid(row=1, column=1, sticky="ew", pady=4, padx=(0, 8))
        self._bind_entry_focus(self.latitude_entry)

        ttk.Label(self, text="Longitude:").grid(row=1, column=2, sticky="w")
        self.longitude_var = tk.StringVar()
        self.longitude_entry = ttk.Entry(self, textvariable=self.longitude_var, width=18)
        self.longitude_entry.grid(row=1, column=3, sticky="ew", pady=4)
        self._bind_entry_focus(self.longitude_entry)

        ttk.Label(self, text="Analysis Name:").grid(row=2, column=0, sticky="w")
        self.analysis_name_var = tk.StringVar()
        self.analysis_name_entry = ttk.Entry(self, textvariable=self.analysis_name_var)
        self.analysis_name_entry.grid(row=2, column=1, columnspan=3, sticky="ew", pady=4)
        self._bind_entry_focus(self.analysis_name_entry)

        ttk.Label(self, text="Scenarios").grid(row=3, column=0, sticky="w")
        ttk.Label(self, text="Models").grid(row=3, column=2, sticky="w", padx=(8, 0))

        self._all_scenarios = self._future_scenarios()
        self.scen_dd = MultiSelectDropdown(
            self,
            placeholder="Select scenarios",
            width=40,
            command=self._refresh_models,
        )
        self.scen_dd.grid(row=4, column=0, rowspan=2, sticky="ew", padx=(0, 4))
        self.scen_dd.set_items(self._all_scenarios)

        self.model_dd = MultiSelectDropdown(self, placeholder="Select models", width=40)
        self.model_dd.grid(row=4, column=2, rowspan=2, sticky="ew", padx=(8, 4))

        yr_frame = ttk.LabelFrame(self, text="Year Ranges", padding=6)
        yr_frame.grid(row=7, column=0, columnspan=4, sticky="ew", pady=(8, 4))

        baseline_start_frame = ttk.Frame(yr_frame)
        baseline_start_frame.grid(row=0, column=0, sticky="w", padx=(0, 20))
        ttk.Label(baseline_start_frame, text="Baseline Start:").pack(side="left")
        self.bl_start_var = tk.StringVar(value="1950")
        self.bl_start_entry = self._build_year_entry(baseline_start_frame, self.bl_start_var)
        self.bl_start_entry.pack(side="left")

        baseline_end_frame = ttk.Frame(yr_frame)
        baseline_end_frame.grid(row=0, column=1, sticky="w", padx=(0, 20))
        ttk.Label(baseline_end_frame, text="End:").pack(side="left")
        self.bl_end_var = tk.StringVar(value="1999")
        self.bl_end_entry = self._build_year_entry(baseline_end_frame, self.bl_end_var)
        self.bl_end_entry.pack(side="left")

        future_start_frame = ttk.Frame(yr_frame)
        future_start_frame.grid(row=0, column=2, sticky="w", padx=(0, 20))
        ttk.Label(future_start_frame, text="Future Start:").pack(side="left")
        self.fut_start_var = tk.StringVar(value="2050")
        self.fut_start_entry = self._build_year_entry(future_start_frame, self.fut_start_var)
        self.fut_start_entry.pack(side="left")

        future_end_frame = ttk.Frame(yr_frame)
        future_end_frame.grid(row=0, column=3, sticky="w")
        ttk.Label(future_end_frame, text="End:").pack(side="left")
        self.fut_end_var = tk.StringVar(value="2099")
        self.fut_end_entry = self._build_year_entry(future_end_frame, self.fut_end_var)
        self.fut_end_entry.pack(side="left")

        pg_frame = ttk.Frame(yr_frame)
        pg_frame.grid(row=1, column=0, columnspan=2, sticky="w", pady=(8, 0))
        ttk.Label(pg_frame, text="PG Reliability:").pack(side="left")
        self.pg_reliability_var = tk.StringVar(value="98%")
        self.pg_reliability_cb = ttk.Combobox(
            pg_frame,
            textvariable=self.pg_reliability_var,
            values=list(PG_RELIABILITY_MAP.keys()),
            width=14,
            height=8,
            state="readonly",
        )
        self.pg_reliability_cb.pack(side="left")

        station_frame = ttk.Frame(yr_frame)
        station_frame.grid(row=1, column=2, columnspan=2, sticky="w", padx=(24, 0), pady=(8, 0))
        ttk.Label(station_frame, text="Station Grid:").pack(side="left")
        self.station_grid_var = tk.StringVar(value="1x1")
        grid_rb_frame = ttk.Frame(station_frame)
        grid_rb_frame.pack(side="left", padx=(6, 0))
        ttk.Radiobutton(grid_rb_frame, text="1×1 (1 station)", variable=self.station_grid_var, value="1x1").pack(side="left", padx=(0, 8))
        ttk.Radiobutton(grid_rb_frame, text="2×2 (4 stations)", variable=self.station_grid_var, value="2x2").pack(side="left", padx=(0, 8))
        ttk.Radiobutton(grid_rb_frame, text="3×3 (9 stations)", variable=self.station_grid_var, value="3x3").pack(side="left")

        param_frame = ttk.LabelFrame(self, text="Index Parameters", padding=6)
        param_frame.grid(row=8, column=0, columnspan=4, sticky="nsew", pady=(4, 4))

        canvas = tk.Canvas(param_frame, height=180, highlightthickness=0)
        scroll = ttk.Scrollbar(param_frame, orient="vertical", command=canvas.yview)
        inner = ttk.Frame(canvas)
        self._param_canvas = canvas
        self._param_inner = inner
        inner.bind("<Configure>", lambda e: canvas.configure(scrollregion=canvas.bbox("all")))
        self._param_canvas_window = canvas.create_window((0, 0), window=inner, anchor="nw")
        canvas.configure(yscrollcommand=scroll.set)
        canvas.bind("<Configure>", self._on_param_canvas_configure)

        canvas.grid(row=0, column=0, columnspan=7, sticky="nsew")
        scroll.grid(row=0, column=7, sticky="ns")
        param_frame.rowconfigure(0, weight=1)
        param_frame.columnconfigure(0, weight=1)
        param_frame.columnconfigure(4, weight=1)
        param_frame.columnconfigure(1, minsize=112)
        param_frame.columnconfigure(2, minsize=112)
        param_frame.columnconfigure(3, minsize=14)
        param_frame.columnconfigure(5, minsize=112)
        param_frame.columnconfigure(6, minsize=112)

        def _on_mousewheel(event):
            canvas.yview_scroll(int(-1 * (event.delta / 120)), "units")

        canvas.bind_all("<MouseWheel>", _on_mousewheel, add="+")

        self.ts_vars = {}
        self.sum_vars = {}
        self._category_param_keys = {}
        self._param_header_frames = []
        self._param_name_labels = []
        self._tooltip = None
        self._param_action_width = 92
        self._param_label_width_px = 220

        grouped_params = []
        current_category = None
        current_items = []

        for key, display, ptype, unit, category in PARAMETER_REGISTRY:
            if category != current_category:
                if current_category is not None:
                    grouped_params.append((current_category, current_items))
                current_category = category
                current_items = []
            current_items.append((key, display, unit))
        if current_category is not None:
            grouped_params.append((current_category, current_items))

        midpoint = (len(grouped_params) + 1) // 2
        left_groups = grouped_params[:midpoint]
        right_groups = grouped_params[midpoint:]

        inner.columnconfigure(0, weight=1, minsize=330, uniform="param_labels")
        inner.columnconfigure(1, minsize=112, uniform="param_ts")
        inner.columnconfigure(2, minsize=112, uniform="param_sum")
        inner.columnconfigure(3, minsize=14)
        inner.columnconfigure(4, weight=1, minsize=330, uniform="param_labels")
        inner.columnconfigure(5, minsize=112, uniform="param_ts")
        inner.columnconfigure(6, minsize=112, uniform="param_sum")

        self._build_header_row(inner)
        self._build_param_column(inner, left_groups, start_col=0, start_row=1)
        self._build_param_column(inner, right_groups, start_col=4, start_row=1)

        sel_frame = ttk.Frame(param_frame)
        sel_frame.grid(row=1, column=0, columnspan=7, sticky="w", pady=(4, 0))
        ttk.Button(sel_frame, text="Select All", command=self._select_all_params).pack(side="left", padx=(0, 8))
        ttk.Button(sel_frame, text="Deselect All", command=self._deselect_all_params).pack(side="left")

        self.run_btn = ttk.Button(self, text="Calculate Index Parameters", command=self._run_calculation)
        self.run_btn.grid(row=10, column=0, columnspan=3, sticky="ew", pady=(8, 4))
        self.stop_btn = ttk.Button(self, text="Stop", command=self._on_stop, state="disabled")
        self.stop_btn.grid(row=10, column=3, sticky="ew", pady=(8, 4), padx=(8, 0))

        self.pb = ttk.Progressbar(self, mode="determinate", maximum=1)
        self.pb.grid(row=11, column=0, columnspan=4, sticky="ew")
        self.pb_label = ttk.Label(self, text="")
        self.pb_label.grid(row=12, column=0, columnspan=4, sticky="w")

        self.log = tk.Text(self, height=6, wrap="word", state="disabled")
        self.log.grid(row=13, column=0, columnspan=4, sticky="nsew", pady=(4, 0))

        self.columnconfigure(0, weight=0, minsize=110)
        self.columnconfigure(1, weight=1)
        self.columnconfigure(2, weight=0, minsize=110)
        self.columnconfigure(3, weight=1)
        self.rowconfigure(8, weight=1)
        self.rowconfigure(13, weight=1)

        self._refresh_models()
        self._us_grid_station_catalog = None
        self._job_id = None
        self._job_log_offset = 0
        self._job_notified = False
        self._job_summary_shown = False
        for col in range(4):
            yr_frame.columnconfigure(col, weight=1, uniform="year_pairs")
        self._apply_param_layout(self.winfo_width() or 920)

    def _build_year_entry(self, parent, variable):
        entry = tk.Entry(
            parent,
            textvariable=variable,
            width=8,
            state="normal",
            justify="left",
            relief="solid",
            borderwidth=1,
        )
        self._bind_entry_focus(entry)
        return entry

    def _bind_entry_focus(self, entry):
        entry.bind("<Button-1>", lambda event, widget=entry: self._focus_entry(widget), add="+")
        entry.bind("<ButtonRelease-1>", lambda event, widget=entry: self._focus_entry(widget), add="+")

    def _focus_entry(self, entry):
        try:
            entry.focus_force()
            if hasattr(entry, "icursor"):
                entry.icursor("end")
        except Exception:
            pass

    def refresh_dropdowns(self):
        self._all_scenarios = self._future_scenarios()
        self.scen_dd.set_items(self._all_scenarios, selected_values=self.scen_dd.get_selected_values())
        self._refresh_models()

    def _future_scenarios(self):
        return [name for name in list_subdirs(self.c.scenario_root) if name.lower() != HISTORICAL_SCENARIO]

    def _refresh_models(self, event=None):
        scenarios = self.scen_dd.get_selected_values() or self._all_scenarios
        models = sorted({model for scenario in scenarios for model in list_subdirs(self.c.scenario_root / scenario)})

        previous = set(self.model_dd.get_selected_values())

        items = [{"label": model, "value": model} for model in models]

        valid_previous = [
            model for model in previous
            if any(item.get("value") == model and item.get("selectable", True) for item in items)
        ]
        self.model_dd.set_items(items, selected_values=valid_previous)

    def _on_param_canvas_configure(self, event):
        try:
            event.widget.itemconfigure(self._param_canvas_window, width=event.width)
            self._apply_param_layout(event.width)
        except Exception:
            pass

    def _apply_param_layout(self, width):
        if not hasattr(self, "_param_inner") or self._param_inner is None:
            return

        usable_width = max(640, int(width) - 24)
        half_width = max(260, usable_width // 2)
        gap_width = 14
        action_width = max(78, min(104, (half_width - 180) // 2))
        label_width = max(150, half_width - (action_width * 2) - gap_width)

        self._param_action_width = action_width
        self._param_label_width_px = label_width

        self._param_inner.columnconfigure(0, minsize=label_width, weight=1, uniform="param_labels")
        self._param_inner.columnconfigure(1, minsize=action_width, uniform="param_ts")
        self._param_inner.columnconfigure(2, minsize=action_width, uniform="param_sum")
        self._param_inner.columnconfigure(4, minsize=label_width, weight=1, uniform="param_labels")
        self._param_inner.columnconfigure(5, minsize=action_width, uniform="param_ts")
        self._param_inner.columnconfigure(6, minsize=action_width, uniform="param_sum")

        for header_frame in self._param_header_frames:
            try:
                header_frame.columnconfigure(0, minsize=label_width, weight=1)
                header_frame.columnconfigure(1, minsize=action_width)
                header_frame.columnconfigure(2, minsize=action_width)
            except Exception:
                pass

        approx_chars = max(12, label_width // 8)
        for label, full_text in self._param_name_labels:
            display_text = self._truncate_text(full_text, approx_chars)
            try:
                label.configure(text=display_text, width=approx_chars)
            except Exception:
                pass

    def _truncate_text(self, text, max_chars):
        text = str(text)
        if len(text) <= max_chars:
            return text
        return text[: max(1, max_chars - 3)].rstrip() + "..."

    def _bind_tooltip(self, widget, text):
        widget.bind("<Enter>", lambda event, value=text: self._show_tooltip(event, value), add="+")
        widget.bind("<Leave>", self._hide_tooltip, add="+")

    def _show_tooltip(self, event, text):
        if not text:
            return
        self._hide_tooltip()
        self._tooltip = tk.Toplevel(self)
        self._tooltip.wm_overrideredirect(True)
        x = event.widget.winfo_rootx() + 12
        y = event.widget.winfo_rooty() + event.widget.winfo_height() + 4
        self._tooltip.wm_geometry(f"+{x}+{y}")
        label = tk.Label(
            self._tooltip,
            text=text,
            justify="left",
            background="#fffde8",
            relief="solid",
            borderwidth=1,
            padx=6,
            pady=4,
        )
        label.pack()

    def _hide_tooltip(self, event=None):
        if self._tooltip is not None:
            try:
                self._tooltip.destroy()
            except Exception:
                pass
            self._tooltip = None

    def _build_header_row(self, parent):
        heading_bg = "#dfe6e0"
        heading_fg = "#266b66"
        headers = [
            (0, "Parameter", "w", (6, 8)),
            (4, "Parameter", "w", (6, 8)),
        ]
        for col, text, anchor, padx in headers:
            label = tk.Label(
                parent,
                text=text,
                bg=heading_bg,
                fg=heading_fg,
                font=("Segoe UI", 9, "bold"),
                anchor=anchor,
                padx=6,
                pady=5,
            )
            label.grid(row=0, column=col, sticky="ew", padx=padx, pady=(0, 6))

    def _create_param_checkbox(self, parent, variable):
        bg = ttk.Style().lookup("TFrame", "background") or "#eef0ed"
        return tk.Checkbutton(
            parent,
            variable=variable,
            onvalue=True,
            offvalue=False,
            bg=bg,
            activebackground=bg,
            highlightthickness=0,
            bd=0,
            relief="flat",
            padx=0,
            pady=0,
        )

    def _grid_centered_cell_widget(self, parent, row, column, widget_factory):
        cell = ttk.Frame(parent)
        cell.grid(row=row, column=column, sticky="ew")
        cell.columnconfigure(0, weight=1)
        cell.rowconfigure(0, weight=1)
        widget = widget_factory(cell)
        widget.grid(row=0, column=0)

    def _build_param_column(self, parent, groups, start_col, start_row):
        row_idx = start_row
        for category, items in groups:
            category_keys = [key for key, _, unit in items]
            self._category_param_keys[category] = category_keys

            header_frame = ttk.Frame(parent)
            header_frame.grid(row=row_idx, column=start_col, columnspan=3, sticky="ew", pady=(6, 2), padx=(6, 0))
            header_frame.columnconfigure(0, weight=1, minsize=330)
            header_frame.columnconfigure(1, minsize=112)
            header_frame.columnconfigure(2, minsize=112)
            self._param_header_frames.append(header_frame)

            ttk.Label(
                header_frame,
                text=f"-- {category} --",
                font=("Segoe UI", 8, "bold"),
                foreground="#5A6C7D",
            ).grid(row=0, column=0, sticky="w")
            ttk.Button(
                header_frame,
                text="Time Series",
                style="Secondary.TButton",
                command=lambda cat=category: self._toggle_category_column(cat, "time_series"),
            ).grid(row=0, column=1, sticky="ew", padx=(8, 4))
            ttk.Button(
                header_frame,
                text="Summary",
                style="Secondary.TButton",
                command=lambda cat=category: self._toggle_category_column(cat, "summary"),
            ).grid(row=0, column=2, sticky="ew")
            row_idx += 1

            for key, display, unit in items:
                full_text = f"{display} ({unit})"
                label = tk.Label(
                    parent,
                    text=full_text,
                    anchor="w",
                    justify="left",
                    background=ttk.Style().lookup("TFrame", "background") or "#eef0ed",
                    foreground=ttk.Style().lookup("TLabel", "foreground") or "#266b66",
                    font=("Segoe UI", 8),
                )
                label.grid(row=row_idx, column=start_col, sticky="w", padx=(6, 8))
                self._param_name_labels.append((label, full_text))
                self._bind_tooltip(label, full_text)

                is_derived = key in (
                    "pg_change_factor_high",
                    "pg_change_factor_low",
                    "final_projected_pg_high",
                    "final_projected_pg_low",
                )

                ts_var = tk.BooleanVar(value=False)
                sum_var = tk.BooleanVar(value=False)
                self.ts_vars[key] = ts_var
                self.sum_vars[key] = sum_var

                if not is_derived:
                    self._grid_centered_cell_widget(
                        parent,
                        row_idx,
                        start_col + 1,
                        lambda cell, variable=ts_var: self._create_param_checkbox(cell, variable),
                    )
                else:
                    self._grid_centered_cell_widget(
                        parent,
                        row_idx,
                        start_col + 1,
                        lambda cell: ttk.Label(cell, text="-"),
                    )
                self._grid_centered_cell_widget(
                    parent,
                    row_idx,
                    start_col + 2,
                    lambda cell, variable=sum_var: self._create_param_checkbox(cell, variable),
                )
                row_idx += 1

    def _toggle_category_column(self, category, column):
        keys = self._category_param_keys.get(category, [])
        if not keys:
            return

        if column == "time_series":
            selectable_keys = [key for key in keys if self._supports_time_series(key)]
            target_vars = [self.ts_vars[key] for key in selectable_keys]
        elif column == "summary":
            selectable_keys = list(keys)
            target_vars = [self.sum_vars[key] for key in selectable_keys]
        else:
            return

        if not target_vars:
            return

        next_state = not all(var.get() for var in target_vars)
        for var in target_vars:
            var.set(next_state)

    def _supports_time_series(self, key):
        return key not in (
            "pg_change_factor_high",
            "pg_change_factor_low",
            "final_projected_pg_high",
            "final_projected_pg_low",
        )

    def _clear_category_params(self, category):
        for key in self._category_param_keys.get(category, []):
            self.ts_vars[key].set(False)
            self.sum_vars[key].set(False)

    def _select_all_params(self):
        for key, value in self.ts_vars.items():
            value.set(self._supports_time_series(key))
        for value in self.sum_vars.values():
            value.set(True)

    def _deselect_all_params(self):
        for value in self.ts_vars.values():
            value.set(False)
        for value in self.sum_vars.values():
            value.set(False)

    def _append_log(self, text):
        self.log.configure(state="normal")
        self.log.insert(tk.END, text + "\n")
        self.log.see(tk.END)
        self.log.configure(state="disabled")

    def _station_catalog(self):
        if self._us_grid_station_catalog is not None:
            return self._us_grid_station_catalog

        workbook_path = Path(getattr(self.c, "us_grid_list_path", ""))
        if workbook_path.exists():
            stations_by_id = {}
            workbook = load_workbook(workbook_path, read_only=True, data_only=True)
            try:
                sheet = workbook.active
                for row in sheet.iter_rows(values_only=True):
                    if not row or len(row) < 3:
                        continue
                    station = str(row[0] or "").strip()
                    try:
                        station_lat = float(row[1])
                        station_lon = float(row[2])
                    except (TypeError, ValueError):
                        continue
                    if not station:
                        continue
                    stations_by_id[station] = {
                        "station": station,
                        "station_lat": station_lat,
                        "station_lon": station_lon,
                    }
            finally:
                workbook.close()

            self._us_grid_station_catalog = list(stations_by_id.values())
            return self._us_grid_station_catalog

        stations_by_id = {}
        for city_info in getattr(self.c, "city_to_info", {}).values():
            if not isinstance(city_info, dict):
                continue

            station = str(city_info.get("station", "")).strip()
            station_lat = city_info.get("station_lat")
            station_lon = city_info.get("station_lon")
            if station and station_lat is not None and station_lon is not None:
                stations_by_id[station] = {
                    "station": station,
                    "station_lat": float(station_lat),
                    "station_lon": float(station_lon),
                }

            for item in city_info.get("grid_stations") or []:
                station = str(item.get("station", "")).strip()
                station_lat = item.get("station_lat")
                station_lon = item.get("station_lon")
                if station and station_lat is not None and station_lon is not None:
                    stations_by_id[station] = {
                        "station": station,
                        "station_lat": float(station_lat),
                        "station_lon": float(station_lon),
                    }

        self._us_grid_station_catalog = list(stations_by_id.values())
        return self._us_grid_station_catalog

    def _haversine_miles(self, lat1, lon1, lat2, lon2):
        radius_mi = 3958.7613
        lat1_rad = math.radians(lat1)
        lon1_rad = math.radians(lon1)
        lat2_rad = math.radians(lat2)
        lon2_rad = math.radians(lon2)
        dlat = lat2_rad - lat1_rad
        dlon = lon2_rad - lon1_rad
        a = math.sin(dlat / 2) ** 2 + math.cos(lat1_rad) * math.cos(lat2_rad) * math.sin(dlon / 2) ** 2
        return 2 * radius_mi * math.asin(math.sqrt(a))

    def _euclidean_distance(self, lat1, lon1, lat2, lon2):
        return ((lat2 - lat1) ** 2 + (lon2 - lon1) ** 2) ** 0.5

    def _station_count_for_grid(self, station_grid):
        return {"1x1": 1, "2x2": 4, "3x3": 9}.get(station_grid, 1)

    def _normalize_longitude(self, longitude):
        station_catalog = self._station_catalog()
        if not station_catalog:
            return longitude, False

        negative_lons = sum(1 for item in station_catalog if float(item["station_lon"]) < 0.0)
        catalog_is_mostly_western = negative_lons >= max(1, int(len(station_catalog) * 0.8))
        if catalog_is_mostly_western and longitude > 0.0:
            return -longitude, True
        return longitude, False

    def _nearest_model_grid_points(self, latitude, longitude, station_grid, scenarios, models, baseline_year):
        if not scenarios or not models:
            return []

        scenario = HISTORICAL_SCENARIO
        model = models[0]
        model_base = Path(self.c.scenario_root) / scenario / model
        tasmax_path, _, _ = resolve_paths_for_year(baseline_year, model_base, model, scenario)
        if not Path(tasmax_path).exists():
            return []

        nearest_points = nearest_grid_points_from_file(
            tasmax_path,
            target_lon=longitude,
            target_lat=latitude,
            count=self._station_count_for_grid(station_grid),
        )

        results = []
        for item in nearest_points:
            station_lat = float(item["lat"])
            station_lon = float(item["lon"])
            distance_mi = self._haversine_miles(latitude, longitude, station_lat, station_lon)
            results.append(
                {
                    "station": f"GRID_{station_lat:.6f}_{station_lon:.6f}",
                    "station_lat": station_lat,
                    "station_lon": station_lon,
                    "distance_mi": distance_mi,
                }
            )
        return results

    def _nearest_stations(self, latitude, longitude, station_grid, scenarios=None, models=None, baseline_year=None):
        if scenarios and models and baseline_year is not None:
            try:
                model_grid_points = self._nearest_model_grid_points(
                    latitude,
                    longitude,
                    station_grid,
                    scenarios,
                    models,
                    baseline_year,
                )
                if model_grid_points:
                    return model_grid_points
            except Exception:
                pass

        station_catalog = self._station_catalog()
        if not station_catalog:
            raise ValueError("No station coordinate data is available in the configured closest-stations CSV.")

        ranked = []
        for item in station_catalog:
            rank_distance = self._euclidean_distance(
                latitude,
                longitude,
                float(item["station_lat"]),
                float(item["station_lon"]),
            )
            ranked.append(
                {
                    "station": item["station"],
                    "station_lat": float(item["station_lat"]),
                    "station_lon": float(item["station_lon"]),
                    "rank_distance": rank_distance,
                    "distance_mi": rank_distance,
                }
            )

        ranked.sort(key=lambda item: item["rank_distance"])
        return ranked[: self._station_count_for_grid(station_grid)]

    def _format_location_label(self, latitude, longitude):
        return f"Lat {latitude:.6f}, Lon {longitude:.6f}"

    def _validate_inputs(self):
        analysis_name = self.analysis_name_var.get().strip()
        if not analysis_name:
            messagebox.showerror("Missing", "Enter an analysis name.")
            return None

        try:
            latitude = float(self.latitude_var.get().strip())
            longitude = float(self.longitude_var.get().strip())
        except ValueError:
            messagebox.showerror("Invalid", "Latitude and longitude must be numeric.")
            return None

        if not (-90.0 <= latitude <= 90.0):
            messagebox.showerror("Invalid", "Latitude must be between -90 and 90.")
            return None

        if not (-180.0 <= longitude <= 180.0):
            messagebox.showerror("Invalid", "Longitude must be between -180 and 180.")
            return None

        longitude, longitude_corrected = self._normalize_longitude(longitude)

        scenarios = self.scen_dd.get_selected_values()
        if not scenarios:
            messagebox.showerror("Missing", "Select at least one scenario.")
            return None

        models = self.model_dd.get_selected_values()
        if not models:
            messagebox.showerror("Missing", "Select at least one model.")
            return None

        try:
            bl_start = int(self.bl_start_var.get())
            bl_end = int(self.bl_end_var.get())
            fut_start = int(self.fut_start_var.get())
            fut_end = int(self.fut_end_var.get())
        except ValueError:
            messagebox.showerror("Invalid", "Year values must be numeric.")
            return None

        if bl_start > bl_end or fut_start > fut_end:
            messagebox.showerror("Invalid", "Start year must be <= end year.")
            return None

        if not (BASELINE_YEAR_MIN <= bl_start <= BASELINE_YEAR_MAX and BASELINE_YEAR_MIN <= bl_end <= BASELINE_YEAR_MAX):
            messagebox.showerror(
                "Invalid Baseline Period",
                f"Baseline years must stay within {BASELINE_YEAR_MIN}–{BASELINE_YEAR_MAX}.",
            )
            return None

        if not (FUTURE_YEAR_MIN <= fut_start <= FUTURE_YEAR_MAX and FUTURE_YEAR_MIN <= fut_end <= FUTURE_YEAR_MAX):
            messagebox.showerror(
                "Invalid Future Period",
                f"Future years must stay within {FUTURE_YEAR_MIN}–{FUTURE_YEAR_MAX}.",
            )
            return None

        baseline_years = bl_end - bl_start + 1
        future_years = fut_end - fut_start + 1

        if baseline_years < 30:
            messagebox.showerror(
                "Invalid Baseline Period",
                "Baseline year range must be at least 30 years.",
            )
            return None

        selected = [key for key in self.ts_vars if self.ts_vars[key].get() or self.sum_vars[key].get()]
        if not selected:
            messagebox.showerror("Missing", "Select at least one parameter.")
            return None

        if future_years < 30:
            if any(key in PRECIP_PARAM_KEYS for key in selected):
                messagebox.showerror(
                    "Invalid Future Period",
                    "Future period is less than 30 years. Precipitation analysis is disabled for <30 years, and 50 years is recommended.",
                )
                return None

            proceed = messagebox.askyesno(
                "Future Period Warning",
                "Using less than 30 years for future analysis is not recommended. Choosing 5 years may result in greater variability.\n\nDo you want to continue?",
            )
            if not proceed:
                return None

        reliability_label = self.pg_reliability_var.get().strip()
        if reliability_label not in PG_RELIABILITY_MAP:
            messagebox.showerror("Invalid", "Select a valid PG reliability option.")
            return None

        ts_flags = {key: self.ts_vars[key].get() for key in self.ts_vars}
        sum_flags = {key: self.sum_vars[key].get() for key in self.sum_vars}

        return {
            "analysis_name": analysis_name,
            "latitude": latitude,
            "longitude": longitude,
            "longitude_corrected": longitude_corrected,
            "location_label": self._format_location_label(latitude, longitude),
            "scenarios": scenarios,
            "models": models,
            "bl_start": bl_start,
            "bl_end": bl_end,
            "fut_start": fut_start,
            "fut_end": fut_end,
            "selected": selected,
            "ts_flags": ts_flags,
            "sum_flags": sum_flags,
            "reliability_z": PG_RELIABILITY_MAP[reliability_label],
            "reliability_label": reliability_label,
            "station_grid": self.station_grid_var.get().strip(),
        }

    def _run_calculation(self):
        inputs = self._validate_inputs()
        if inputs is None:
            return

        try:
            selected_stations = self._nearest_stations(
                inputs["latitude"],
                inputs["longitude"],
                inputs["station_grid"],
                inputs["scenarios"],
                inputs["models"],
                inputs["bl_start"],
            )
        except Exception as exc:
            messagebox.showerror("Station Lookup Error", str(exc))
            return

        if not selected_stations:
            messagebox.showerror("Station Lookup Error", "No nearby stations were found for the entered coordinates.")
            return

        inputs["selected_stations"] = selected_stations

        analysis_root_base = Path(self.c.output_base_index_analysis)
        analysis_root_base.mkdir(parents=True, exist_ok=True)
        analysis_dir, renamed = self._prepare_analysis_dir(analysis_root_base, inputs["analysis_name"])
        inputs["analysis_dir"] = analysis_dir

        if renamed:
            messagebox.showinfo(
                "Analysis Folder Renamed",
                f"Duplicate analysis folder found.\n\nNew folder name:\n{analysis_dir.name}",
            )

        self._write_analysis_options(inputs, analysis_dir)

        per_location_steps = len(selected_stations) * (
            len(inputs["models"]) + len(inputs["scenarios"]) * len(inputs["models"]) * 2
        )
        total_steps = per_location_steps
        self.pb.configure(maximum=total_steps, value=0)
        self.pb_label.config(text=f"0 / {total_steps} steps")

        self.log.configure(state="normal")
        self.log.delete("1.0", tk.END)
        self.log.configure(state="disabled")

        def runner(ctx):
            ctx.log("Starting index parameter calculation...")
            if inputs.get("longitude_corrected"):
                ctx.log(f"Corrected longitude to western hemisphere: {inputs['longitude']:.6f}")
            ctx.log(f"Analysis output folder: {analysis_dir}")
            summary_payloads = []

            location_info = {
                "lat": inputs["latitude"],
                "lon": inputs["longitude"],
                "grid_stations": selected_stations,
            }

            def progress_cb(step, total, msg):
                ctx.set_progress(step, total_steps, f"{inputs['location_label']}: {msg}")

            try:
                summary, ts_files = run_index_calculation(
                    city=inputs["location_label"],
                    city_info=location_info,
                    scenario_root=self.c.scenario_root,
                    hcd_input_dir=self.c.hcd_input_dir,
                    scenarios=inputs["scenarios"],
                    models=inputs["models"],
                    baseline_start=inputs["bl_start"],
                    baseline_end=inputs["bl_end"],
                    future_start=inputs["fut_start"],
                    future_end=inputs["fut_end"],
                    selected_params=inputs["selected"],
                    ts_flags=inputs["ts_flags"],
                    summary_flags=inputs["sum_flags"],
                    output_root=str(analysis_dir),
                    progress_callback=progress_cb,
                    reliability_z=inputs["reliability_z"],
                    station_grid=inputs["station_grid"],
                    cancel_event=ctx.cancel_event,
                )

                if ts_files:
                    ctx.log(f"{inputs['location_label']}: saved {len(ts_files)} time-series CSV files.")
                    for path in ts_files[:5]:
                        ctx.log(f"  -> {path}")
                    if len(ts_files) > 5:
                        ctx.log(f"  ... and {len(ts_files) - 5} more.")

                sum_params = [key for key in inputs["selected"] if inputs["sum_flags"].get(key, False)]
                if sum_params:
                    exported_files = self._export_summary_csvs(
                        summary_results=summary,
                        location_name=inputs["location_label"],
                        location_info=location_info,
                        scenarios=inputs["scenarios"],
                        station_grid=inputs["station_grid"],
                        analysis_dir=analysis_dir,
                    )
                    for exported_file in exported_files:
                        ctx.log(f"Summary CSV saved: {exported_file}")
                    summary_payloads.append((inputs["location_label"], location_info, summary))

                if summary_payloads:
                    self._write_root_summary_files(summary_payloads, analysis_dir)
                else:
                    self._write_root_summary_files([], analysis_dir)
                    ctx.log("No summary parameters selected. Done.")

                ctx.log("Calculation complete.")
                return {
                    "summary": summary,
                    "location_label": inputs["location_label"],
                    "has_summary": bool(summary_payloads),
                    "analysis_dir": str(analysis_dir),
                }
            except OperationCancelledError:
                shutil.rmtree(analysis_dir, ignore_errors=True)
                raise

        self._job_id = self.c.job_manager.submit(
            title=f"Index Analysis: {inputs['location_label']}",
            kind="index-analysis",
            runner=runner,
        )
        self._job_log_offset = 0
        self._job_notified = False
        self._job_summary_shown = False
        self.stop_btn.config(state="normal")
        self.after(100, self._poll_job)

    def _prepare_analysis_dir(self, parent_dir, analysis_name):
        base_name = self._sanitize_name(analysis_name)
        candidate = Path(parent_dir) / base_name
        if not candidate.exists():
            candidate.mkdir(parents=True, exist_ok=False)
            return candidate, False

        suffix = 1
        while True:
            renamed = Path(parent_dir) / f"{base_name} ({suffix})"
            if not renamed.exists():
                renamed.mkdir(parents=True, exist_ok=False)
                return renamed, True
            suffix += 1

    def _sanitize_name(self, value):
        cleaned = "".join(ch for ch in value.strip() if ch not in '<>:"/\\|?*')
        cleaned = " ".join(cleaned.split())
        return cleaned or "Analysis"

    def _write_analysis_options(self, inputs, analysis_dir):
        param_lookup = {key: (display, unit) for key, display, _ptype, unit, _category in PARAMETER_REGISTRY}

        def selected_lines(flag_map):
            lines = []
            for key in inputs["selected"]:
                if flag_map.get(key, False):
                    display, unit = param_lookup[key]
                    lines.append(f"- {display} ({unit})")
            return lines or ["- None"]

        lines = [
            f"Analysis Name: {inputs['analysis_name']}",
            f"Output Folder: {analysis_dir}",
            f"Input Latitude: {inputs['latitude']}",
            f"Input Longitude: {inputs['longitude']}",
            f"Scenarios: {', '.join(inputs['scenarios'])}",
            f"Models: {', '.join(inputs['models'])}",
            f"Baseline Years: {inputs['bl_start']} - {inputs['bl_end']}",
            f"Future Years: {inputs['fut_start']} - {inputs['fut_end']}",
            f"PG Reliability: {inputs['reliability_label']}",
            f"Station Grid: {inputs['station_grid']}",
            "Selected Stations:",
        ]
        lines.extend(
            f"- {item['station']} | lat={item['station_lat']:.6f} | lon={item['station_lon']:.6f} | distance_mi={item['distance_mi']:.3f}"
            for item in inputs.get("selected_stations", [])
        )
        lines.extend([
            "",
            "Time Series Parameters:",
        ])
        lines.extend(selected_lines(inputs["ts_flags"]))
        lines.extend(["", "Summary Parameters:"])
        lines.extend(selected_lines(inputs["sum_flags"]))

        (Path(analysis_dir) / "analysis_options.txt").write_text("\n".join(lines) + "\n", encoding="utf-8")

    def _export_summary_csvs(self, summary_results, location_name, location_info, scenarios, station_grid, analysis_dir):
        exported_files = []
        lon = (location_info or {}).get("lon", "")
        lat = (location_info or {}).get("lat", "")
        location_safe = self._sanitize_name(location_name).replace(" ", "_")
        selected_stations = next(
            (
                pdata.get("stations", [])
                for pdata in summary_results.values()
                if pdata.get("stations")
            ),
            [],
        )
        station_desc = "; ".join(
            f"{item.get('station')} ({item.get('lat')}, {item.get('lon')} | {float(item.get('distance_mi', 0.0)):.3f} mi)"
            for item in selected_stations
        )

        for scenario in scenarios:
            model_names = self._scenario_models(summary_results, scenario)
            if not model_names:
                continue

            scenario_dir = Path(analysis_dir) / self._sanitize_name(scenario)
            scenario_dir.mkdir(parents=True, exist_ok=True)
            csv_path = scenario_dir / f"{location_safe}_{station_grid}.csv"

            with csv_path.open("w", newline="", encoding="utf-8") as handle:
                writer = csv.writer(handle)
                writer.writerow([
                    "Scenario", scenario,
                    "Location", location_name,
                    "Stations", station_desc,
                    "Input Longitude", lon,
                    "Input Latitude", lat,
                    "Grid", station_grid,
                ])

                header = ["Section", "Category", "Parameter", "Unit"]
                for model_name in model_names:
                    header.extend([
                        f"{model_name} Historical Value",
                        f"{model_name} Future Value",
                        f"{model_name} Change Factor",
                        f"{model_name} Normalized Delta (deg C)",
                        f"{model_name} % Difference",
                    ])
                writer.writerow(header)

                for param_key, display_name, _ptype, unit, category in PARAMETER_REGISTRY:
                    pdata = summary_results.get(param_key)
                    if not pdata:
                        continue

                    row = [
                        pdata.get("section", "CMIP Tool Parameters"),
                        category,
                        display_name,
                        self._display_unit_for_csv(unit),
                    ]

                    for model_name in model_names:
                        label = f"{scenario}/{model_name}"
                        mdata = pdata.get("models", {}).get(label, {})
                        historical_value = mdata.get("historical_model")
                        fm = mdata.get("future_model")
                        diff = self._summary_change_factor(param_key, historical_value, fm)
                        pct = mdata.get("pct_change")
                        row.extend([
                            self._format_summary_value(historical_value),
                            self._format_summary_value(fm),
                            self._format_summary_value(diff),
                            self._format_summary_value(mdata.get("normalized_delta_c")),
                            f"{pct:.2f}" if pct is not None else "",
                        ])

                    writer.writerow(row)

            exported_files.append(str(csv_path))

        return exported_files

    def _write_root_summary_files(self, summary_payloads, analysis_dir):
        csv_path = Path(analysis_dir) / "summary.csv"
        txt_path = Path(analysis_dir) / "summary.txt"

        with csv_path.open("w", newline="", encoding="utf-8") as csv_handle, txt_path.open("w", encoding="utf-8") as txt_handle:
            writer = csv.writer(csv_handle)
            writer.writerow([
                "Location",
                "Input Latitude",
                "Input Longitude",
                "Stations",
                "Section",
                "Scenario",
                "Parameter",
                "Unit",
                "Model",
                "Historical Model Value",
                "Future Model Value",
                "Change Factor",
                "Normalized Delta (deg C)",
                "% Difference",
            ])

            for location_name, location_info, summary_results in summary_payloads:
                selected_stations = next(
                    (
                        pdata.get("stations", [])
                        for pdata in summary_results.values()
                        if pdata.get("stations")
                    ),
                    [],
                )
                station_desc = ", ".join(
                    f"{item.get('station')} [{item.get('lat')}, {item.get('lon')}; {float(item.get('distance_mi', 0.0)):.3f} mi]"
                    for item in selected_stations
                )
                txt_handle.write(f"Location: {location_name}\n")
                txt_handle.write(f"Input Latitude: {(location_info or {}).get('lat', '')}\n")
                txt_handle.write(f"Input Longitude: {(location_info or {}).get('lon', '')}\n")
                txt_handle.write(f"Stations: {station_desc}\n")

                scenarios = sorted(
                    {
                        label.split("/", 1)[0]
                        for pdata in summary_results.values()
                        for label in pdata.get("models", {})
                    }
                )

                for scenario in scenarios:
                    txt_handle.write(f"Scenario: {scenario}\n")
                    for param_key, display_name, _ptype, unit, _category in PARAMETER_REGISTRY:
                        pdata = summary_results.get(param_key)
                        if not pdata:
                            continue

                        for model_name in self._scenario_models(summary_results, scenario):
                            label = f"{scenario}/{model_name}"
                            mdata = pdata.get("models", {}).get(label, {})
                            historical_value = mdata.get("historical_model")
                            future_value = mdata.get("future_model")
                            hist_str = self._format_summary_value(historical_value)
                            aggregate_str = self._format_summary_value(mdata.get("future_model"))
                            diff_str = self._format_summary_value(
                                self._summary_change_factor(param_key, historical_value, future_value)
                            )
                            pct_str = self._format_percent_value(mdata.get("pct_change"))
                            writer.writerow([
                                location_name,
                                (location_info or {}).get("lat", ""),
                                (location_info or {}).get("lon", ""),
                                station_desc,
                                pdata.get("section", "CMIP Tool Parameters"),
                                scenario,
                                display_name,
                                self._display_unit_for_csv(unit),
                                model_name,
                                hist_str,
                                aggregate_str,
                                diff_str,
                                self._format_summary_value(mdata.get("normalized_delta_c")),
                                pct_str,
                            ])
                            txt_handle.write(
                                f"- {display_name} [{self._display_unit_for_csv(unit)}] | {model_name} | historical={hist_str} | aggregate={aggregate_str} | change_factor={diff_str} | pct_difference={pct_str}\n"
                            )
                    txt_handle.write("\n")

        if summary_payloads:
            self._write_extended_summary_workbook(summary_payloads, analysis_dir)

    def _write_extended_summary_workbook(self, summary_payloads, analysis_dir):
        workbook = Workbook()
        workbook.remove(workbook.active)

        cmip_ws = workbook.create_sheet(title="CMIP_Parameters")
        pg_ws = workbook.create_sheet(title="Pavement_PG_Parameters")
        validation_ws = workbook.create_sheet(title="Validation_Missing_Extra")
        validation_ws.append(["Location", "Validation Type", "Parameter"])

        for location_name, location_info, summary_results in summary_payloads:
            validation = summary_results.get("__validation__", {})
            for missing_name in validation.get("missing", []):
                validation_ws.append([location_name, "Missing", missing_name])
            for extra_name in validation.get("extra", []):
                validation_ws.append([location_name, "Extra", extra_name])

            station_summaries = next(
                (
                    pdata.get("station_summaries", [])
                    for key, _display_name, _ptype, _unit, _category in PARAMETER_REGISTRY
                    for pdata in [summary_results.get(key)]
                    if pdata and pdata.get("station_summaries")
                ),
                [],
            )
            model_labels = sorted(
                {
                    label
                    for pdata in summary_results.values()
                    if isinstance(pdata, dict)
                    for label in pdata.get("models", {})
                }
            )

            self._append_extended_summary_section(
                cmip_ws,
                location_name,
                location_info,
                summary_results,
                station_summaries,
                model_labels,
                include_pg=False,
            )
            self._append_extended_summary_section(
                pg_ws,
                location_name,
                location_info,
                summary_results,
                station_summaries,
                model_labels,
                include_pg=True,
            )

        wrap_alignment = Alignment(wrap_text=True, vertical="top")
        header_fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")  # Light gray
        header_font = Font(bold=True, size=11)
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin")
        )
        
        for worksheet in (cmip_ws, pg_ws):
            for row in worksheet.iter_rows():
                for cell in row:
                    cell.alignment = wrap_alignment
            
            # Apply header formatting to first 2 rows
            for row_idx in range(1, 3):
                for col_idx in range(1, worksheet.max_column + 1):
                    cell = worksheet.cell(row=row_idx, column=col_idx)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.border = thin_border
            
            # Freeze panes: freeze columns A-D and rows 1-2
            worksheet.freeze_panes = "E3"

        workbook_path = Path(analysis_dir) / "extended_summary.xlsx"
        workbook.save(workbook_path)

    def _append_extended_summary_section(
        self,
        worksheet,
        location_name,
        location_info,
        summary_results,
        station_summaries,
        model_labels,
        include_pg,
    ):
        if worksheet.max_row > 1:
            worksheet.append([])

        start_row = worksheet.max_row + 1
        block_width = (3 * len(model_labels)) + 3

        # Parse scenarios and models from labels (do this once at the beginning)
        scenario_model_map = {}
        for label in model_labels:
            if "/" in label:
                scenario, model = label.split("/", 1)
            else:
                scenario = "Unknown"
                model = label
            if scenario not in scenario_model_map:
                scenario_model_map[scenario] = []
            scenario_model_map[scenario].append((model, label))

        # Header Row 1: Station info header
        header_row_1 = ["", "", "", ""]  # Empty base columns for alignment
        for station_index, station_summary in enumerate(station_summaries, start=1):
            header_row_1.extend(
                [self._station_block_title(station_index, station_summary)]
                + [""] * (block_width - 1)
            )
            if station_index < len(station_summaries):
                header_row_1.append("")
        worksheet.append(header_row_1)
        worksheet.row_dimensions[start_row].height = 90

        # Header Row 2: Combined column labels (Scenario - Model - DataType)
        header_row_2 = ["Location", "Category", "Parameter", "Unit"]
        for station_index, _station_summary in enumerate(station_summaries, start=1):
            for scenario in sorted(scenario_model_map.keys()):
                models = scenario_model_map[scenario]
                for model, _label in models:
                    # Format: "RCP 4.5 -ACCESS1-0 Historical" etc.
                    header_row_2.extend([
                        f"{scenario} -{model} Historical",
                        f"{scenario} -{model} Future",
                        f"{scenario} -{model} Change Factor"
                    ])
            header_row_2.extend(["Ensemble Historical", "Ensemble Future", "Ensemble Change Factor"])
            if station_index < len(station_summaries):
                header_row_2.append("")
        worksheet.append(header_row_2)
        worksheet.row_dimensions[start_row + 1].height = 30

        # Do NOT merge cells - keep all columns independent and unmerged

        worksheet.column_dimensions["A"].width = 24
        worksheet.column_dimensions["B"].width = 24
        worksheet.column_dimensions["C"].width = 30
        worksheet.column_dimensions["D"].width = 10
        for station_index in range(len(station_summaries)):
            base_col = 5 + station_index * (block_width + 1)
            for label_index, _label in enumerate(model_labels):
                hist_col = base_col + (label_index * 3)
                future_col = hist_col + 1
                diff_col = hist_col + 2
                worksheet.column_dimensions[self._excel_column_name(hist_col)].width = 16
                worksheet.column_dimensions[self._excel_column_name(future_col)].width = 16
                worksheet.column_dimensions[self._excel_column_name(diff_col)].width = 18
            worksheet.column_dimensions[self._excel_column_name(base_col + block_width - 3)].width = 16
            worksheet.column_dimensions[self._excel_column_name(base_col + block_width - 2)].width = 14
            worksheet.column_dimensions[self._excel_column_name(base_col + block_width - 1)].width = 16
            if station_index < len(station_summaries) - 1:
                worksheet.column_dimensions[self._excel_column_name(base_col + block_width)].width = 3

        for param_key, display_name, _ptype, unit, category in PARAMETER_REGISTRY:
            if include_pg != (param_key in PG_PARAMETER_KEYS):
                continue

            pdata = summary_results.get(param_key)
            if not pdata:
                continue

            row = [
                self._format_location_text(location_name, location_info),
                category,
                display_name,
                self._display_unit_for_excel_summary(include_pg),
            ]
            param_station_summaries = pdata.get("station_summaries", [])
            for station_index in range(len(station_summaries)):
                station_summary = param_station_summaries[station_index] if station_index < len(param_station_summaries) else {}

                # Append data for each scenario (in sorted order)
                for scenario in sorted(scenario_model_map.keys()):
                    models = scenario_model_map[scenario]
                    for model, label in models:
                        model_payload = station_summary.get("models", {}).get(label, {})
                        station_hist = model_payload.get("historical")
                        station_value = model_payload.get("future")
                        station_diff = self._summary_change_factor(param_key, station_hist, station_value)
                        row.append(self._format_summary_value(station_hist))
                        row.append(self._format_summary_value(station_value))
                        row.append(self._format_summary_value(station_diff))
                
                # Append Ensemble value and diff
                ensemble_value = station_summary.get("ensemble_average")
                ensemble_hist_values = [
                    model_payload.get("historical")
                    for model_payload in station_summary.get("models", {}).values()
                    if isinstance(model_payload.get("historical"), (int, float))
                ]
                ensemble_hist = (
                    sum(ensemble_hist_values) / len(ensemble_hist_values)
                    if ensemble_hist_values else None
                )
                ensemble_diff = self._summary_change_factor(param_key, ensemble_hist, ensemble_value)
                row.append(self._format_summary_value(ensemble_hist))
                row.append(self._format_summary_value(ensemble_value))
                row.append(self._format_summary_value(ensemble_diff))
                if station_index < len(station_summaries) - 1:
                    row.append("")
            worksheet.append(row)

    def _station_block_title(self, station_index, station_summary):
        return (
            f"Station - {station_index} ID:\n{station_summary.get('station', '')}\n"
            f"Lat/Lon:\n{self._format_summary_value(station_summary.get('lat'))}, "
            f"{self._format_summary_value(station_summary.get('lon'))}\n"
            f"Distance (mi):\n{self._format_summary_value(station_summary.get('distance_mi'))}"
        )

    def _format_location_text(self, location_name, location_info):
        lat = (location_info or {}).get("lat", "")
        lon = (location_info or {}).get("lon", "")
        if lat != "" and lon != "":
            return f"Lat {lat}, Lon {lon}"
        return str(location_name)

    def _display_unit_for_excel_summary(self, include_pg):
        return "°C" if include_pg else "°F"

    def _excel_column_name(self, column_number):
        name = ""
        while column_number > 0:
            column_number, remainder = divmod(column_number - 1, 26)
            name = chr(65 + remainder) + name
        return name

    def _display_unit_for_excel_summary(self, include_pg):
        return "\u00B0C" if include_pg else "\u00B0F"

    def _display_unit_for_csv(self, unit):
        return str(unit).replace("\u00B0F", "deg F").replace("\u00B0C", "deg C")

    def _excel_scalar_value(self, value):
        if isinstance(value, (int, float)):
            return float(value)
        return None

    def _summary_change_factor(self, param_key, historic_value, future_value):
        if not isinstance(historic_value, (int, float)) or not isinstance(future_value, (int, float)):
            return None
        parameter_type = PARAMETER_METADATA_BY_KEY.get(param_key, {}).get("parameter_type")
        if parameter_type == "precip":
            if historic_value == 0:
                return None
            return future_value / historic_value
        return future_value - historic_value

    def _format_summary_value(self, value):
        if value is None:
            return "Not available"
        return f"{value:.4f}" if isinstance(value, (int, float)) else str(value)

    def _format_percent_value(self, value):
        if value is None:
            return ""
        return f"{value:.2f}" if isinstance(value, (int, float)) else str(value)

    def _scenario_models(self, summary_results, scenario):
        prefix = f"{scenario}/"
        model_names = set()
        for pdata in summary_results.values():
            for label in pdata.get("models", {}):
                if label.startswith(prefix):
                    model_names.add(label[len(prefix):])
        return sorted(model_names)

    def _on_stop(self):
        if self._job_id is not None:
            self.c.job_manager.cancel(self._job_id)
            self.stop_btn.config(state="disabled")
            self.pb_label.config(text="Stopping...")

    def _poll_job(self):
        if self._job_id is None:
            return

        snapshot = self.c.job_manager.get_job_snapshot(self._job_id)
        if snapshot is None:
            return

        new_logs = snapshot["logs"][self._job_log_offset:]
        self._job_log_offset = len(snapshot["logs"])
        for line in new_logs:
            self._append_log(line)

        self.pb["maximum"] = max(1, int(snapshot["progress_total"]))
        self.pb["value"] = min(int(snapshot["progress_current"]), int(self.pb["maximum"]))
        self.pb_label.config(text=f"{snapshot['progress_current']} / {snapshot['progress_total']} steps - {snapshot['status']}")

        if snapshot["status"] in {"queued", "running", "cancelling"}:
            self.stop_btn.config(state="normal")
            self.after(200, self._poll_job)
            return

        self.stop_btn.config(state="disabled")

        result = snapshot.get("result")
        if (
            snapshot["status"] == "completed"
            and not self._job_summary_shown
            and self.c.current_frame_name == "IndexParameterFrame"
            and isinstance(result, dict)
            and result.get("has_summary")
        ):
            self._show_summary_popup(result["summary"], city_name=result["location_label"])
            self._job_summary_shown = True

        if not self._job_notified and self.c.current_frame_name == "IndexParameterFrame":
            if snapshot["status"] == "cancelled":
                messagebox.showinfo("Calculation Stopped", "Index parameter calculation was stopped.")
            elif snapshot["status"] == "failed":
                messagebox.showerror("Calculation Failed", snapshot["error"] or "Unknown error")
            elif snapshot["status"] == "completed" and not (isinstance(result, dict) and result.get("has_summary")):
                messagebox.showinfo("Calculation Complete", "Index parameter calculation completed.")
            self._job_notified = True

    def _show_summary_popup(self, summary_results, city_name=None):
        popup = tk.Toplevel(self)
        popup.title("Index Parameter Summary" if not city_name else f"Index Parameter Summary - {city_name}")
        popup.geometry("1100x650")
        popup.transient(self)
        popup.grab_set()

        title_text = "Index Parameter Summary" if not city_name else f"Index Parameter Summary - {city_name}"
        ttk.Label(popup, text=title_text, font=("Segoe UI", 14, "bold")).pack(pady=(10, 5))

        tree_frame = ttk.Frame(popup)
        tree_frame.pack(fill="both", expand=True, padx=10, pady=5)

        all_models = set()
        for pdata in summary_results.values():
            all_models.update(pdata.get("models", {}).keys())
        all_models = sorted(all_models)

        columns = ["unit", "historic"]
        col_headers = {"unit": "Unit", "historic": "Historical Baseline"}
        for model in all_models:
            safe_model = model.replace("/", "_")
            for suffix, label in [("_fm", "Aggregate"), ("_diff", "Difference"), ("_pct", "% Difference")]:
                cid = safe_model + suffix
                columns.append(cid)
                col_headers[cid] = f"{model} {label}"

        tree = ttk.Treeview(tree_frame, columns=columns, show="tree headings", height=20)
        tree.heading("#0", text="Parameter")
        tree.column("#0", width=200, minwidth=150)
        for column in columns:
            tree.heading(column, text=col_headers[column])
            tree.column(column, width=100, minwidth=70, anchor="center")

        vsb = ttk.Scrollbar(tree_frame, orient="vertical", command=tree.yview)
        hsb = ttk.Scrollbar(tree_frame, orient="horizontal", command=tree.xview)
        tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)
        tree.grid(row=0, column=0, sticky="nsew")
        vsb.grid(row=0, column=1, sticky="ns")
        hsb.grid(row=1, column=0, sticky="ew")
        tree_frame.rowconfigure(0, weight=1)
        tree_frame.columnconfigure(0, weight=1)

        current_cat = None
        for param_key, display_name, ptype, unit, category in PARAMETER_REGISTRY:
            if param_key not in summary_results:
                continue

            pdata = summary_results[param_key]

            if category != current_cat:
                current_cat = category
                tree.insert("", "end", iid=f"cat_{category}", text=f"-- {category} --", open=True)

            parent = f"cat_{category}"
            hist_val = pdata.get("historic")

            if isinstance(hist_val, dict):
                hist_str = "; ".join(f"{k}: {v:.2f}" for k, v in hist_val.items() if v is not None)
            elif hist_val is not None:
                hist_str = f"{hist_val:.4f}"
            else:
                hist_str = "N/A"

            vals = {"unit": unit, "historic": hist_str}

            for model in all_models:
                safe_model = model.replace("/", "_")
                mdata = pdata.get("models", {}).get(model, {})

                fm = mdata.get("future_model")
                diff = mdata.get("difference")
                pct = mdata.get("pct_change")
                if isinstance(fm, dict) or isinstance(diff, dict):
                    fm_str = "; ".join(f"{k}: {v:.2f}" for k, v in (fm or {}).items() if v is not None)
                    vals[safe_model + "_fm"] = fm_str
                    vals[safe_model + "_diff"] = "; ".join(f"{k}: {v:.2f}" for k, v in (diff or {}).items() if v is not None)
                    vals[safe_model + "_pct"] = "-"
                else:
                    vals[safe_model + "_fm"] = f"{fm:.4f}" if fm is not None else "N/A"
                    vals[safe_model + "_diff"] = f"{diff:.4f}" if diff is not None else "N/A"
                    vals[safe_model + "_pct"] = f"{pct:.2f}%" if pct is not None else "N/A"

            tree.insert(parent, "end", text=display_name, values=[vals.get(column, "") for column in columns])

        def _export():
            path = filedialog.asksaveasfilename(
                title="Export Summary as CSV",
                defaultextension=".csv",
                filetypes=[("CSV", "*.csv"), ("All files", "*")],
                parent=popup,
            )
            if not path:
                return

            with open(path, "w", newline="", encoding="utf-8") as handle:
                writer = csv.writer(handle)
                header = ["Category", "Parameter", "Unit", "Historical Baseline"]
                for model in all_models:
                    header.extend([f"{model} Aggregate", f"{model} Difference", f"{model} % Difference"])
                writer.writerow(header)

                for param_key, display_name, ptype, unit, category in PARAMETER_REGISTRY:
                    if param_key not in summary_results:
                        continue

                    pdata = summary_results[param_key]
                    hist_val = pdata.get("historic")

                    if isinstance(hist_val, dict):
                        hist_str = "; ".join(f"{k}: {v:.4f}" for k, v in hist_val.items() if v is not None)
                    elif hist_val is not None:
                        hist_str = f"{hist_val:.4f}"
                    else:
                        hist_str = ""

                    row = [category, display_name, unit, hist_str]

                    for model in all_models:
                        mdata = pdata.get("models", {}).get(model, {})
                        fm = mdata.get("future_model")
                        diff = mdata.get("difference")
                        pct = mdata.get("pct_change")
                        if isinstance(fm, dict) or isinstance(diff, dict):
                            row.extend([
                                "; ".join(f"{k}: {v:.4f}" for k, v in (fm or {}).items() if v is not None),
                                "; ".join(f"{k}: {v:.4f}" for k, v in (diff or {}).items() if v is not None),
                                "",
                            ])
                        else:
                            row.extend([
                                f"{fm:.4f}" if fm is not None else "",
                                f"{diff:.4f}" if diff is not None else "",
                                f"{pct:.2f}" if pct is not None else "",
                            ])
                    writer.writerow(row)

            messagebox.showinfo("Export", f"Summary exported to:\n{path}", parent=popup)

        btn_frame = ttk.Frame(popup)
        btn_frame.pack(fill="x", padx=10, pady=(0, 10))
        ttk.Button(btn_frame, text="Export as CSV", command=_export).pack(side="left")
        ttk.Button(btn_frame, text="Close", command=popup.destroy).pack(side="right")
